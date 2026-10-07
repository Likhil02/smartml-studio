"""Dynamic Excel experiment parser (two-row merged headers, comma decimals, ranges, text flags)."""
import re
from io import BytesIO
import openpyxl


def _num(v):
    if v is None: return None
    if isinstance(v, (int, float)): return float(v)
    try: return float(str(v).strip().replace(",", "."))
    except ValueError: return None


def _range(v):
    """'10-100' -> (min, max, mid, None); 'Overfitting' -> (None, None, None, 'Overfitting')."""
    if v is None: return None, None, None, None
    if isinstance(v, (int, float)): return float(v), float(v), float(v), None
    s = str(v).strip().replace(",", ".")
    m = re.fullmatch(r"(\d+(?:\.\d+)?)\s*-\s*(\d+(?:\.\d+)?)", s)
    if m:
        a, b = float(m[1]), float(m[2]); return a, b, (a + b) / 2, None
    n = _num(s)
    return (n, n, n, None) if n is not None else (None, None, None, s)


def _header_row(ws):
    for r in range(1, min(ws.max_row, 10) + 1):
        vals = [str(c.value or "").lower() for c in ws[r]]
        if any(v.startswith("sample") for v in vals) and any("epoch" in v for v in vals):
            return r
    return None


def _columns(ws, r1):
    top = [c.value for c in ws[r1]]
    sub = [c.value for c in ws[r1 + 1]] if r1 < ws.max_row else []
    has_sub = any(isinstance(x, str) and x.strip().lower() in ("acc", "tes acc", "loss", "tes loss") for x in sub)
    cols, parent = {}, ""
    for i, t in enumerate(top):
        if t is not None: parent = str(t).strip().lower()
        elif not has_sub: continue
        s = str(sub[i]).strip().lower() if has_sub and i < len(sub) and sub[i] is not None else ""
        key = None
        if parent.startswith("accuracy"): key = "test_acc" if s.startswith("tes") else "acc"
        elif parent.startswith("loss"): key = "test_loss" if s.startswith("tes") else "loss"
        elif parent.startswith("sample"): key = "sample"
        elif parent.startswith("epoch"): key = "epoch"
        elif "batch" in parent: key = "batch_size"
        elif "learning" in parent: key = "learning_rate"
        elif "original" in parent: key = "original_acc"
        elif "different" in parent: key = "different_acc"
        if key and key not in cols: cols[key] = i
    return cols, has_sub


def parse_workbook(source):
    try:
        wb = openpyxl.load_workbook(BytesIO(source) if isinstance(source, (bytes, bytearray)) else source, data_only=True)
    except Exception as e:
        raise ValueError(f"Not a valid .xlsx workbook: {e}")
    rows, notes, sheets, warnings = [], [], [], []
    for ws in wb:
        r1 = _header_row(ws)
        if r1 is None:
            if ws.title.lower().startswith("note"):
                for r in ws.iter_rows(values_only=True):
                    notes.extend(v.strip() for v in r if isinstance(v, str) and len(v.strip()) > 15)
            continue
        cols, has_sub = _columns(ws, r1)
        missing = {"sample", "epoch", "batch_size", "learning_rate", "acc"} - set(cols)
        if missing:
            warnings.append(f"Sheet '{ws.title}' skipped: missing {sorted(missing)}"); continue
        sheets.append(ws.title)
        for r in ws.iter_rows(min_row=r1 + 1 + (1 if has_sub else 0), values_only=True):
            g = lambda k: r[cols[k]] if k in cols and cols[k] < len(r) else None
            smp, ep, bs, lr = _num(g("sample")), _num(g("epoch")), _num(g("batch_size")), _num(g("learning_rate"))
            if None in (smp, ep, bs, lr): continue
            oa, da = _range(g("original_acc")), _range(g("different_acc"))
            rows.append({"sheet": ws.title, "sample": int(smp), "epoch": int(ep), "batch_size": int(bs), "learning_rate": lr,
                         "acc": _num(g("acc")), "test_acc": _num(g("test_acc")), "loss": _num(g("loss")), "test_loss": _num(g("test_loss")),
                         "original_min": oa[0], "original_max": oa[1], "original_mid": oa[2], "original_flag": oa[3],
                         "different_min": da[0], "different_max": da[1], "different_mid": da[2], "different_flag": da[3]})
    if not rows:
        raise ValueError("No experiment rows found. Expected sheets with Sample / Epoch / Batch Size / Learning Rate / Accuracy headers.")
    return {"sheets": sheets, "rows": rows, "notes": notes, "warnings": warnings}


def _avg(xs):
    xs = [x for x in xs if x is not None]
    return sum(xs) / len(xs) if xs else None


def _best_by(rows, key):
    g = {}
    for r in rows:
        if r["test_acc"] is not None: g.setdefault(r[key], []).append(r["test_acc"])
    if not g: return None
    k = max(g, key=lambda x: _avg(g[x]))
    return {"value": k, "avg_test_acc": _avg(g[k]), "n": len(g[k])}


def summarize(rows):
    valid = [r for r in rows if r["test_acc"] is not None]
    key = lambda r: (r["test_acc"], -(r["test_loss"] if r["test_loss"] is not None else 9e9), r["different_mid"] or 0)
    grp = {}
    for r in valid: grp.setdefault((r["sample"], r["epoch"], r["batch_size"], r["learning_rate"]), []).append(r["test_acc"])
    varied = [{"sample": k[0], "epoch": k[1], "batch_size": k[2], "learning_rate": k[3], "runs": len(v), "min": min(v), "max": max(v)}
              for k, v in grp.items() if len(v) > 1 and max(v) - min(v) > 0.05]
    return {"count": len(rows), "best": max(valid, key=key) if valid else None, "worst": min(valid, key=key) if valid else None,
            "avg_acc": _avg(r["acc"] for r in rows), "avg_test_acc": _avg(r["test_acc"] for r in rows),
            "avg_loss": _avg(r["loss"] for r in rows), "best_learning_rate": _best_by(rows, "learning_rate"),
            "best_batch_size": _best_by(rows, "batch_size"), "best_epoch": _best_by(rows, "epoch"),
            "flagged_rows": sum(1 for r in rows if r["original_flag"] or r["different_flag"]),
            "repeated_runs_with_variation": varied}


def chart_series(rows):
    def agg(x, y):
        g = {}
        for r in rows:
            if r.get(y) is not None: g.setdefault(r[x], []).append(r[y])
        return [{"x": k, "y": round(_avg(v), 4), "n": len(v)} for k, v in sorted(g.items())]
    return {"sample_vs_acc": agg("sample", "test_acc"), "epoch_vs_acc": agg("epoch", "test_acc"),
            "batch_vs_acc": agg("batch_size", "test_acc"), "lr_vs_acc": agg("learning_rate", "test_acc"),
            "epoch_vs_loss": agg("epoch", "test_loss"), "lr_vs_loss": agg("learning_rate", "test_loss"),
            "sample_vs_different": agg("sample", "different_mid"),
            "original_vs_different": [{"original": r["original_mid"], "different": r["different_mid"], "sample": r["sample"]}
                                      for r in rows if r["original_mid"] is not None and r["different_mid"] is not None]}


def filter_rows(rows, **f):
    for k, v in f.items():
        if v is not None: rows = [r for r in rows if r[k] == v]
    return rows
